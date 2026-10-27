"""Read-only XLSX/CSV contract. Never execute source formulas, instructions or code."""
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import PureWindowsPath
import csv
import hashlib
import re
import zipfile
from .engine import TABLES


@dataclass
class FinancialData:
    tables: dict
    label: str = "OWNER PROVIDED / UNVERIFIED"

    @property
    def raw_counts(self):
        return {k: len(self.tables.get(k, [])) for k in TABLES}


def _value(value):
    from datetime import date, datetime
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and re.match(r"^\s*(?:[=+@]|-[A-Za-z])", value):
        raise ValueError("Không nhận công thức/code trong ô nguồn; cần giá trị gốc, không dùng cache")
    if isinstance(value, str) and (len(value) > 500 or any(ord(c) < 32 and c not in "\t\n\r" for c in value)):
        raise ValueError("Ô văn bản quá dài hoặc có ký tự điều khiển")
    return value


def _rows(table, records, filename, digest):
    headers, rows = None, []
    for physical, values in records:
        if not values or all(v is None or str(v).strip() == "" for v in values):
            continue
        if headers is None:
            headers = [str(v).strip() for v in values]
            if len(set(headers)) != len(headers) or set(headers) != set(TABLES[table]):
                raise ValueError(f"{table}: header phải đúng data dictionary, không thêm cột chỉ dẫn/credential")
            continue
        if len(values) != len(headers):
            raise ValueError(f"{table} dòng {physical}: số cột khác header")
        row = {key: _value(value) for key, value in zip(headers, values)}
        row["_source"] = f"{filename}#{digest[:12]}::{table}!row {physical}"
        rows.append(row)
        if len(rows) > 100000 or len(rows)*len(headers) > 1000000:
            raise ValueError("Tối đa 100.000 dòng/một triệu ô mỗi bảng")
    if headers is None:
        raise ValueError(table+": thiếu header")
    return rows


def _read_financial(data, name, table=""):
    filename = PureWindowsPath(name).name
    digest = hashlib.sha256(data).hexdigest()
    if len(data) > 20*1024*1024:
        raise ValueError("File vượt 20 MB")
    tables = {}
    if filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 64*1024*1024 or any("vbaproject" in e.filename.lower() or "externallink" in e.filename.lower() for e in entries):
                raise ValueError("Workbook quá lớn hoặc chứa macro/external link")
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if any(sheet not in TABLES for sheet in book.sheetnames) or not book.sheetnames:
                raise ValueError("Chỉ nhận các sheet canonical Chief Auditor")
            if sum((s.max_row or 0)*(s.max_column or 0) for s in book) > 1000000:
                raise ValueError("Workbook vượt một triệu ô")
            for sheet in book:
                tables[sheet.title] = _rows(sheet.title, enumerate(sheet.values, 1), filename, digest)
                if sum(len(rows)*len(TABLES[k]) for k, rows in tables.items()) > 1000000:
                    raise ValueError("Workbook vượt một triệu ô dữ liệu")
        finally:
            book.close()
    elif filename.lower().endswith(".csv"):
        table = table or PureWindowsPath(filename).stem
        if table not in TABLES:
            raise ValueError("CSV cần tên bảng canonical: orders.csv, trips.csv…")
        stream = csv.reader(StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
        def records():
            while True:
                start = stream.line_num+1
                try:
                    values = next(stream)
                except StopIteration:
                    return
                end = stream.line_num
                yield str(start) if start == end else f"{start}:{end}", values
        tables[table] = _rows(table, records(), filename, digest)
    else:
        raise ValueError("Chỉ nhận .xlsx/.csv UTF-8")
    return FinancialData(tables, "SYNTHETIC DEMO / UNVERIFIED" if filename.lower().startswith("synthetic_") else "OWNER PROVIDED / UNVERIFIED")


def read_financial(data, name, table=""):
    try:
        return _read_financial(data, name, table)
    except ValueError:
        raise
    except Exception:
        raise ValueError("Không đọc được workbook/CSV; kiểm tra file không hỏng, không mã hóa và đúng UTF-8/schema.") from None


def read_many(files):
    tables, labels = {}, []
    for name, data in files:
        loaded = read_financial(data, name)
        if set(tables) & set(loaded.tables):
            raise ValueError("Trùng bảng giữa các file; cần gộp/đối soát trước")
        tables.update(loaded.tables)
        labels.append(loaded.label)
    return FinancialData(tables, "SYNTHETIC DEMO / UNVERIFIED" if labels and all("SYNTHETIC" in l for l in labels) else "OWNER PROVIDED / UNVERIFIED")
