"""Read-only adapters. Workbook formula caches never certify actuals or money."""
from datetime import datetime, date, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET
import hashlib
import csv
import re
import openpyxl
from .domain import number, identity, require_scope, CATEGORIES

ASSUMPTIONS = {5: "days", 6: "kg_day", 7: "pcs_per_kg", 8: "price_pcs", 9: "asis_kg", 10: "benchmark_kg",
    11: "dc_kg", 12: "hub_kg", 13: "last_mile_kg", 14: "labor_kg", 15: "truck8_trip",
    16: "truck8_trips", 17: "truck5_month", 18: "capacity19", 19: "capacity5", 20: "capacity8",
    21: "ncc_package", 22: "ncc_package_drops", 23: "ncc_drops", 24: "surcharge_pcs", 26: "sla_target", 27: "extra_day"}


def guarded_zip(content):
    if len(content) > 20_000_000:
        raise ValueError("File vượt 20 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 100_000_000:
                raise ValueError("Nội dung giải nén vượt 100 MB.")
    except BadZipFile:
        raise ValueError("File Office hỏng hoặc không đúng định dạng .xlsx/.docx.") from None


def source_record(content, filename):
    return {"file": Path(filename).name, "sha256": hashlib.sha256(content).hexdigest(),
            "updated_at": None, "imported_at": datetime.now(timezone.utc).isoformat(), "verification": "UNVERIFIED"}


def read_workbook(content, filename):
    guarded_zip(content)
    source = source_record(content, filename)
    w = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=False, keep_links=False)
    required = {"01_GIA_DINH", "02_PNL_TUYEN", "03_DOI_XE", "04_PNL_DIEM", "05_NHAT_KY_14N", "09_NHAT_KY_XE", "10_DIEM_14N"}
    if not required <= set(w.sheetnames):
        raise ValueError("Workbook không đúng mẫu Supra; thiếu sheet bắt buộc.")
    if sum(s.max_row*s.max_column for s in w) > 500000:
        raise ValueError("Workbook vượt giới hạn ô đọc.")
    unit_checks = {5: "ngày", 6: "kg/ngày", 7: "pcs/kg", 8: "VND/pcs", 15: "VND/chuyến", 17: "VND/tháng", 18: "kg/chuyến", 19: "kg/chuyến", 20: "kg/chuyến"}
    for row, unit in unit_checks.items():
        if w["01_GIA_DINH"][f"C{row}"].value != unit:
            raise ValueError(f"Đơn vị 01_GIA_DINH!C{row} không đúng schema {unit}; không tự đổi kg/tấn/pcs.")
    if w["02_PNL_TUYEN"]["E4"].value != "KG/NGÀY" or w["03_DOI_XE"]["C4"].value != "TẢI/CHUYẾN KG":
        raise ValueError("Tiêu đề/đơn vị tuyến hoặc đội xe không đúng schema.")
    refs = {}
    def get(sheet, cell, numeric=True, optional=False):
        value = w[sheet][cell].value
        if w[sheet][cell].data_type == "f":
            raise ValueError(f"Ô đầu vào {sheet}!{cell} là công thức; cần số nguồn rõ, không dùng cached value.")
        refs[f"{sheet}!{cell}"] = f"{source['file']}#{source['sha256'][:12]}::{sheet}!{cell}"
        return number(value, optional=optional) if numeric else value
    assumptions = {name: get("01_GIA_DINH", f"B{row}", optional=name == "ncc_drops") for row, name in ASSUMPTIONS.items()}
    if assumptions["days"] not in (26, 30) or assumptions["pcs_per_kg"] <= 0 or assumptions["sla_target"] > 1:
        raise ValueError("Cần kỳ 26/30 ngày, hệ số dương và SLA từ 0 đến 1.")
    start = get("01_GIA_DINH", "B25", numeric=False)
    if not isinstance(start, (datetime, date)):
        raise ValueError("Ngày pilot phải là ngày Excel, không đoán văn bản.")
    routes = []
    for row in range(5, 17):
        route = {"id": get("02_PNL_TUYEN", f"A{row}", False), "name": get("02_PNL_TUYEN", f"B{row}", False),
            "flow": get("02_PNL_TUYEN", f"C{row}", False), "stores": get("02_PNL_TUYEN", f"D{row}"),
            "kg": get("02_PNL_TUYEN", f"E{row}"), "source": refs[f"02_PNL_TUYEN!E{row}"]}
        if not route["id"] or route["stores"] != route["stores"].to_integral_value() or route["kg"] <= 0:
            raise ValueError("Tuyến cần mã, số cửa hàng nguyên và kg dương.")
        route["hub"] = "VIET_TRI" if route["flow"] == "HUB VIỆT TRÌ" else "THANH_BA" if route["flow"] == "HUB THANH BA" else "DIRECT_DSD"
        routes.append(route)
    if len({r["id"] for r in routes}) != 12 or sum(r["kg"] for r in routes) != assumptions["kg_day"]:
        raise ValueError("Trùng mã tuyến hoặc kg tuyến không khớp tổng đầu vào.")
    fleets = []
    for row in range(5, 12):
        fleets.append({"id": f"F{row}", "name": get("03_DOI_XE", f"A{row}", False),
            **{name: get("03_DOI_XE", f"{col}{row}", optional=True) for name, col in
               (("vehicles", "B"), ("legal_kg", "C"), ("trips", "D"), ("kg", "F"))},
            "legal_capacity_verified": False, "source": refs[f"03_DOI_XE!C{row}"]})
    stores = []
    for row in range(5, 135):
        stores.append({"id": get("04_PNL_DIEM", f"A{row}", False), "route": get("04_PNL_DIEM", f"B{row}", False),
            "real_id": get("04_PNL_DIEM", f"D{row}", False), "name": get("04_PNL_DIEM", f"E{row}", False),
            "kg": get("04_PNL_DIEM", f"G{row}", optional=True),
            "source": refs[f"04_PNL_DIEM!A{row}"]})
    if len({s["id"] for s in stores}) != len(stores) or any(s["route"] not in {r["id"] for r in routes} for s in stores):
        raise ValueError("Khóa điểm giao trùng hoặc tuyến không tồn tại.")
    actuals, issues = [], []
    sheet = w["05_NHAT_KY_14N"]
    for row in range(5, sheet.max_row+1):
        if not any(sheet[f"{col}{row}"].value is not None for col in ("E", "F", "H", "I", "J", "K", "L", "M", "N", "S", "V")):
            continue
        try:
            day = sheet[f"A{row}"].value
            offset = re.fullmatch(r"='01_GIA_DINH'!\$B\$25\+(\d+)", day) if isinstance(day, str) else None
            if offset:
                day = start + timedelta(days=int(offset.group(1)))
            # Source template dates are the only supported input formula.
            if not isinstance(day, (datetime, date)):
                raise ValueError("Ngày ACTUAL không xác định; không dùng cache công thức.")
            item = {"date": day.date().isoformat() if isinstance(day, datetime) else day.isoformat(),
                "route": get("05_NHAT_KY_14N", f"B{row}", False), "kg": get("05_NHAT_KY_14N", f"E{row}", optional=True),
                "pcs": get("05_NHAT_KY_14N", f"F{row}", optional=True), "drops": get("05_NHAT_KY_14N", f"H{row}", optional=True),
                "on_time": get("05_NHAT_KY_14N", f"S{row}", optional=True),
                "costs": {cat: get("05_NHAT_KY_14N", f"{col}{row}", optional=True) for cat, col in zip(CATEGORIES, "IJKLMN")},
                "document_ref": get("05_NHAT_KY_14N", f"V{row}", False), "verification": "UNVERIFIED", "mode": "ACTUAL",
                "source": f"{source['file']}#{source['sha256'][:12]}::05_NHAT_KY_14N!A{row}:V{row}"}
            actuals.append(item)
        except ValueError as exc:
            issues.append({"source": f"05_NHAT_KY_14N!{row}", "error": str(exc)})
    result = {"source": source, "refs": refs, "assumptions": assumptions, "routes": routes, "fleets": fleets,
        "stores": stores, "actuals": actuals, "issues": issues, "pilot_start": start.date().isoformat() if isinstance(start, datetime) else start.isoformat(),
        "mode": "PLAN", "verification": "UNVERIFIED", "customer": "SUPRA", "region": "PHU_THO"}
    result["version"] = identity({"hash": source["sha256"], "adapter": 1})
    return result


def read_document(content, filename):
    guarded_zip(content)
    source = source_record(content, filename)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(BytesIO(content)) as archive:
        body = ET.fromstring(archive.read("word/document.xml")).find("w:body", ns)
    records = []
    for i, node in enumerate(body, 1):
        text = " ".join(t.text or "" for t in node.findall(".//w:t", ns))
        if text:
            records.append({"text": text, "source": f"{source['file']}#{source['sha256'][:12]}::body[{i}]"})
    return {"source": source, "records": records}


def read_actual_csv(content, filename):
    """Canonical route/day ledger. Status and provenance cannot be supplied by CSV."""
    if len(content) > 20_000_000:
        raise ValueError("CSV vượt 20 MB.")
    text = content.decode("utf-8-sig")
    reader = csv.reader(StringIO(text))
    headers = next(reader, None)
    required = {"date", "route", "kg", "pcs", "drops", "on_time", "price_pcs", "document_ref", *CATEGORIES}
    if not headers or len(headers) != len(set(headers)) or not required <= set(headers):
        raise ValueError("Thiếu cột CSV actual chuẩn: " + ", ".join(sorted(required)))
    source = source_record(content, filename)
    records = []
    previous_line = reader.line_num
    for values in reader:
        start_line, end_line = previous_line+1, reader.line_num
        previous_line = end_line
        if not values or not any(values):
            continue
        if len(values) != len(headers):
            raise ValueError("Dòng CSV không đúng số cột ở dòng "+str(start_line))
        row = dict(zip(headers, values))
        require_scope(row)
        if len(records) >= 100000:
            raise ValueError("CSV vượt 100.000 dòng.")
        day = date.fromisoformat(row["date"])
        item = {"date": day.isoformat(), "route": row["route"],
            **{k: number(row.get(k), optional=True, integer=k in ("pcs", "drops", "on_time")) for k in ("kg", "pcs", "drops", "on_time", "price_pcs")},
            "costs": {cat: number(row[cat], optional=True) for cat in CATEGORIES},
            "document_ref": row["document_ref"][:2000], "mode": "ACTUAL", "verification": "UNVERIFIED",
            "source": f"{source['file']}#{source['sha256'][:12]}::CSV:row{start_line}"+(f":{end_line}" if end_line != start_line else "")}
        records.append(item)
    return {"records": records, "source": source}
